"""
Exports the speaker notes out of both decks into readable markdown scripts.

The notes pane inside the .pptx is the source of truth. These markdown files are generated
views of it, so they can never disagree with what is actually in the deck you present from.
Re-run after any change to build_deck.py or build_deck_v3.py.

Usage: python export_scripts.py
"""

import os
import re
import sys

from pptx import Presentation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

OUT_DIR = r"c:\Users\lenovo\EA\projects\ucip"
WPM = 140.0
STAGE_DIRECTION = re.compile(r"\[[^\]]*\]")

SECTIONS = {
    1: "01 · Project overview & participant details",
    2: "02 · Problem understanding",
    3: "03 · Research & insights",
    4: "04 · Proposed solution",
    5: "05 · Prototype / solution design",
    6: "06 · Impact & future scope",
}
V2_SECTIONS = [1, 1, 2, 2, 2, 3, 3, 3, 4, 4, 4, 4, 5, 5, 5, 5, 5, 6, 6, 6]
# v3 carries a 21st slide, the source appendix. It is not presented, so its notes are a single
# bracketed stage direction and it contributes zero spoken words.
V3_SECTIONS = V2_SECTIONS + [6]

DECKS = [
    {
        "pptx": "UCIP-Pitch-Deck.pptx",
        "md": "presentation-script.md",
        "title": "UCIP Presentation Script (v2, technical cut)",
        "sections": V2_SECTIONS,
        "blurb": (
            "Spoken script for `UCIP-Pitch-Deck.pptx`. This is the technical cut: it names PCA, "
            "Kendall tau, NDVI and land-surface temperature directly, and is pitched at a mixed "
            "panel where at least some judges are researchers.\n\n"
            "For the plain-language version of the same deck, see `presentation-script-v3.md`."
        ),
    },
    {
        "pptx": "UCIP-Pitch-Deck-v3.pptx",
        "md": "presentation-script-v3.md",
        "title": "UCIP Presentation Script (v3, plain-language cut)",
        "sections": V3_SECTIONS,
        "blurb": (
            "Spoken script for `UCIP-Pitch-Deck-v3.pptx`. This is the cut to present from: same "
            "layout as v2, no jargon on any slide, written in the first person throughout, and "
            "every slide carries its own source line in the lower band so any claim can be "
            "checked from the slide it appears on. The opening is built around one causal "
            "chain:\n\n"
            "> more heat  →  discomfort for everyone, illness for some, and load on the "
            "hospitals  →  citizens suffer  →  so a city must choose where to act first\n\n"
            "Slide 2 states the chain, and slides 3 to 5 walk it link by link. The fourth link "
            "is the only one you can change, and that is where the tool lives.\n\n"
            "Slide 21 is a source appendix. It is not presented and is not timed; it exists so a "
            "judge can trace any figure without asking.\n\n"
            "For the technical version, see `presentation-script.md`."
        ),
    },
]

TITLE_SKIP = re.compile(r"^\d+\s*/\s*\d+$")


def slide_heading(slide):
    """Human title for a slide: the text set in the largest font.

    Picking the first long string instead would grab screenshot captions on the slides whose
    real title is short, for example "The map".
    """
    best, best_size = None, 0.0
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        txt = " ".join(shape.text_frame.text.split())
        if not txt or txt.isupper() or TITLE_SKIP.match(txt):
            continue
        # Mono is reserved for computed figures (stat boxes, DOIs). A 26 pt "0.978" would
        # otherwise outrank the 25 pt slide heading beside it.
        size = max(
            (r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs
             if r.font.size is not None and r.font.name != "JetBrains Mono"),
            default=0.0,
        )
        if size > best_size:
            best, best_size = txt, size
    return best[:70] if best else "(visual slide)"


def render(deck):
    path = os.path.join(OUT_DIR, deck["pptx"])
    prs = Presentation(path)
    slide_section = deck["sections"]
    if len(slide_section) != len(prs.slides):
        raise ValueError(
            f"{deck['pptx']} has {len(prs.slides)} slides but its section map has "
            f"{len(slide_section)} entries"
        )

    lines = [f"# {deck['title']}", "", deck["blurb"], ""]

    total_words = 0
    rows = []
    bodies = []
    for i, slide in enumerate(prs.slides, start=1):
        notes = ""
        if slide.has_notes_slide:
            notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
        words = len(STAGE_DIRECTION.sub(" ", notes).split())
        total_words += words
        secs = words / WPM * 60
        rows.append((i, slide_section[i - 1], slide_heading(slide), words, secs))
        bodies.append((i, notes, secs))

    total_sec = total_words / WPM * 60
    lines += [
        f"**{len(prs.slides)} slides. About {int(total_sec // 60)} minutes "
        f"{int(total_sec % 60)} seconds of speaking** at 140 words per minute "
        f"({total_words} spoken words).",
        "",
        "Text in square brackets is a delivery cue, not something you say out loud. It is "
        "excluded from the timing above.",
        "",
        "---",
        "",
        "## At a glance",
        "",
        "| # | Section | Slide | Words | Approx. |",
        "|---|---------|-------|-------|---------|",
    ]
    for i, sec, head, words, secs in rows:
        lines.append(f"| {i} | {sec} | {head} | {words} | {int(secs//60)}:{int(secs%60):02d} |")

    lines += ["", "---", "", "## The script", ""]

    last_section = None
    for (i, notes, secs), (_, sec, head, _, _) in zip(bodies, rows):
        if sec != last_section:
            lines += [f"### {SECTIONS[sec]}", ""]
            last_section = sec
        lines += [
            f"#### Slide {i}. {head}",
            "",
            f"*approx. {int(secs//60)}:{int(secs%60):02d}*",
            "",
            notes,
            "",
        ]

    md_path = os.path.join(OUT_DIR, deck["md"])
    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines).rstrip() + "\n")

    print(f"  {deck['md']:<30} {len(prs.slides)} slides, {total_words} words, "
          f"{int(total_sec//60)}:{int(total_sec%60):02d}")


def main():
    print("Exporting speaker notes to markdown:")
    for deck in DECKS:
        render(deck)


if __name__ == "__main__":
    main()
